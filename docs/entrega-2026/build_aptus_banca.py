#!/usr/bin/env python3
"""Revisão editorial orientada pelas duas rubricas e canvas técnico.

As regras gerais e o briefing são conhecidos; o DOCX oficial não foi recuperado.
Não descrever esta reconstrução como reprodução validada do modelo original.
"""
from pathlib import Path
import json
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
VIDEO = ROOT.parent / 'video'
CANVAS = ROOT / 'canvas_premium' / 'APTUS_Canvas_Premium.png'
WORD = ROOT / 'D03_APTUS_Resumo_2026_PREENCHER.docx'

TEXT = {
    'title': 'APTUS: arquitetura de IA para equipes e propostas técnicas',
    'synthesis': 'O Aptus propõe à Nexora um pipeline de IA que encontra profissionais aderentes, apresenta experiências comprovadas e mantém a base atualizada. A contribuição da equipe está na qualidade da implementação proposta: integrar interpretação semântica, regras verificáveis e otimização, preservando a origem de cada decisão. RAG contextualiza o LLM, Jev classifica aderência, Java calcula pontos e OR-Tools compõe equipes sob restrições. O Recovery incorpora provas aceitas e recalcula a proposta. Pendências críticas bloqueiam a confirmação e seguem para revisão. A arquitetura prevê permissões por papel e finalidade antes da busca, parsers isolados, modelos sem privilégios administrativos, auditoria e contingência manual. Dados reais e serviços externos dependem de finalidade, base legal, contrato e avaliação da transferência internacional aplicável. A implantação começa com uma área e mede qualidade, esforço, custo e adesão. O cenário utiliza dados sintéticos; desempenho e viabilidade econômica serão verificados no piloto.',
    'problem': 'A Nexora reúne mais de 1.000 profissionais. O briefing estima 40 disputas anuais e 15–25 horas conjuntas de gestores por disputa: 600–1.000 horas/ano, a validar. Currículos dispersos, desatualização e provas ausentes ocultam competências e prejudicam a composição da equipe. O objetivo é recuperar informações, comprovar sua aderência ao edital e apresentar a equipe disponível, mantendo o acervo atualizado [1].',
    'solution': 'Docling extrai documentos e Qwen estrutura fatos; PostgreSQL guarda fontes e versões. O RAG combina recuperação textual, vetorial e estruturada para contextualizar a análise. Jev avalia correspondências e ambiguidades; Licitações aprova critérios, e Java calcula elegibilidade e pontos. OR-Tools CP-SAT compõe equipes conforme regras e agenda confirmada pelo PMO. Busca SQL preserva candidatos conhecidos. Só provas verificadas pelo RH e aceitas no edital geram pontos após recálculo. O dossiê conserva fatos e citações, com aprovação humana. Para sustentar a adesão, eventos de projetos, campanhas e indicações geram atualizações para confirmação; RH mantém o cadastro [2].',
    'foundation': 'A implementação proposta usa núcleo Java/Spring e PostgreSQL, modelos locais leves, filas limitadas, reuso de verificações válidas e processamento incremental. O custo total inclui implantação amortizada, hardware, APIs, manutenção e horas de todas as áreas. Metas de qualidade, latência e revisão serão pactuadas antes do piloto. A expansão exige benefício líquido positivo, testes de segurança e avaliação jurídica [2,4].',
    'obtained': 'Arquitetura e cenário sintético: 30→31 pontos por composição conjunta; 31→33→35 com dois projetos distintos comprovados (+2 cada), após aceite e recálculo. Mock planejado; integrações, desempenho e ganhos reais ainda não foram validados.',
    'expected': 'Hipótese: 40 disputas/ano × 20 h/disputa, redução líquida de 40% e R$ 160/h liberariam 320 h/ano, equivalentes a R$ 51.200 de capacidade, antes das despesas adicionais. Premissas ilustrativas; economia de caixa não garantida.',
    'partnership': 'A EloGroup atua na combinação de tecnologia, dados e transformação de processos, enquanto o CEFET-MG reúne formação em engenharia de software, inteligência artificial, bancos de dados e pesquisa operacional. Essa afinidade permite desenvolver e avaliar o Aptus como pesquisa aplicada a uma necessidade empresarial concreta. A cooperação proposta começaria com uma área, 50–100 profissionais e 3–5 editais históricos, inicialmente com dados sintéticos e, posteriormente, dados autorizados em modo sombra. A empresa definiria critérios de negócio e disponibilizaria revisores; estudantes e docentes implementariam os módulos, testes e experimentos. O piloto compararia a arquitetura com busca manual e configurações simplificadas, medindo erros, cobertura documental, tempo, custo total, adesão e fila de revisão. RH verificaria evidências, PMO confirmaria disponibilidade e Licitações revisaria critérios e dossiês. A ampliação dependeria de resultados previamente pactuados, avaliação jurídica e testes de segurança. Contratos, finalidade, retenção, direitos dos titulares e eventual transferência internacional seriam tratados antes do uso de dados reais. A continuidade e a propriedade intelectual seguiriam os instrumentos institucionais, com participação da CIE/DEDC. Assim, ensino e pesquisa poderiam produzir uma implementação tecnicamente consistente e economicamente justificável [3,4].',
}

PITCH = [
 ('00:00–00:30','Problema quantificado','Abrir com o caso Nexora e a conta 40 × 15–25 = 600–1.000 horas por ano.',
  'A Nexora reúne mais de mil profissionais. Segundo o briefing, são cerca de quarenta disputas por ano, e os gestores gastam, juntos, de quinze a vinte e cinco horas por disputa procurando e ajustando currículos. Isso representa de seiscentas a mil horas anuais, além do risco de deixar de fora profissionais cuja experiência não foi localizada ou comprovada.'),
 ('00:30–00:55','Contribuição da equipe','Exibir Aptus e o canvas completo, identificado como arquitetura proposta.',
  'A contribuição do Aptus está na qualidade da implementação proposta. Integramos recuperação de evidências, modelos de linguagem, Jev, regras verificáveis e otimização para ligar cada recomendação às suas fontes. Essa arquitetura foi desenhada com responsabilidades claras, custo controlável e revisão humana nos pontos que exigem decisão.'),
 ('00:55–02:00','Pipeline de funcionamento','Percorrer os grupos Evidências, Interpretação e Decisão; destacar a rota SQL independente.',
  'Docling organiza os documentos e o LLM Qwen estrutura as informações. O PostgreSQL preserva fatos, fontes e versões. A recuperação híbrida localiza evidências para contextualizar o modelo pelo RAG. Jev avalia correspondências e ambiguidades entre requisitos e provas.\n\nLicitações revisa os critérios; Java calcula elegibilidade e pontos; o OR-Tools compõe a equipe respeitando papéis e disponibilidade confirmada pelo PMO. A consulta SQL mantém candidatos conhecidos no processo, mesmo quando a similaridade textual é baixa.\n\nNa saída, o LLM redige com fontes, Jev revisa a aderência semântica e Java confere referências. O dossiê passa por aprovação. A interpretação fica ligada a regras reproduzíveis e a responsáveis que podem revisar a decisão.'),
 ('02:00–02:55','Demonstração da lógica','Exibir 30→31; em seguida 31→33→35 com duas provas distintas aceitas e novos cálculos.',
  'O exemplo utiliza dados sintéticos e processamento simulado. Escolher os profissionais isoladamente produz trinta pontos; considerar a composição conjunta chega a trinta e um. O Recovery identifica dois projetos já realizados que ainda precisam de atestados aceitos.\n\nReceber o documento não altera a pontuação. Depois da verificação, do aceite no edital e de novo cálculo, ela passa para trinta e três. Com a segunda prova, chega a trinta e cinco. A interface mostra a regra e a evidência de cada parcela. Essa é uma demonstração da lógica, não uma medição de resultado comercial.'),
 ('02:55–03:30','Viabilidade e impacto','Mostrar arquitetura enxuta, processamento incremental e hipótese 800h × 40% = 320h, com rótulo de projeção.',
  'O plano usa um núcleo Java e PostgreSQL, modelos locais leves e processamento incremental, com filas e limites de execução. O custo total inclui implantação, infraestrutura, APIs, manutenção e revisão humana. No cenário de quarenta disputas de vinte horas, uma redução líquida hipotética de quarenta por cento liberaria trezentas e vinte horas anuais. O piloto precisa demonstrar que o benefício compensa os custos adicionais.'),
 ('03:30–04:10','Confiabilidade segurança e jurídico','Exibir permissões antes da busca, documentos isolados e caminho de revisão humana.',
  'A arquitetura prevê permissões por papel e finalidade antes de recuperar qualquer trecho, parsers isolados e modelos sem privilégios administrativos. Falhas e ambiguidades seguem para revisão, com histórico e contingência manual. O uso de dados reais exige definição de finalidade, base legal e retenção. Serviços externos dependem de avaliação contratual e, quando aplicável, de transferência internacional. Esses controles serão testados antes de ampliar o uso.'),
 ('04:10–04:30','Adesão e implantação','Mostrar confirmação de atualização e as fases Mock, Piloto e Integrações.',
  'Eventos de projetos, campanhas e indicações geram atualizações para confirmação. RH mantém o cadastro, PMO responde pela agenda e Licitações pelos critérios. A implantação começa em uma área, medindo também a adesão e a fila de revisão.'),
 ('04:30–04:50','Cooperação e fechamento','Exibir EloGroup e CEFET-MG como participantes de uma cooperação proposta.',
  'A EloGroup contribuiria com regras e casos de negócio. O CEFET-MG agregaria desenvolvimento de software, inteligência artificial e pesquisa operacional para avaliar o piloto. O Aptus propõe uma implementação moderna, verificável e progressiva, cujo avanço depende de qualidade e viabilidade demonstradas.'),
]


def clean_style(style):
    props = style.element.find(qn('w:pPr'))
    if props is not None:
        for el in list(props):
            if el.tag == qn('w:pBdr'):
                props.remove(el)


def new_document(size=10):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.65)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(.6), Cm(.65)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        s = doc.styles[name]
        s.font.name = 'Arial'
        s.font.color.rgb = RGBColor(15, 19, 26)
        clean_style(s)
    normal = doc.styles['Normal']
    normal.font.size = Pt(size)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.045
    normal.paragraph_format.widow_control = True
    title = doc.styles['Title']
    title.font.size = Pt(14)
    title.font.bold = True
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True
    for name in ('Heading 1', 'Heading 2'):
        s = doc.styles[name]
        s.font.size = Pt(10.5)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(7)
        s.paragraph_format.space_after = Pt(3)
        s.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('APTUS | HackatOhm 2026 | D03 | ')
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(91, 97, 107)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    return doc


def body(doc, text, bold_phrase=None):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_phrase and bold_phrase in text:
        a,b=text.split(bold_phrase,1)
        p.add_run(a)
        p.add_run(bold_phrase).bold=True
        p.add_run(b)
    else:
        p.add_run(text)
    return p


def hyperlink(p, label, url):
    rel=p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'233C73'); pr.append(color)
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'17'); pr.append(size)
    r.append(pr); t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)


def refs(doc):
    doc.add_paragraph('6  Fontes e ferramentas','Heading 1')
    rows=[
      [('[1] EloGroup. Case B Nexora, pp. 14–19. Material do desafio HackatOhm 2026.',None)],
      [('[2] Lewis et al. (2020), RAG','https://ai.meta.com/research/publications/retrieval-augmented-generation-for-knowledge-intensive-nlp-tasks/'),('; ',None),('TypeSafe, Jev','https://docs.typesafe.ai/models'),('; ',None),('Google, OR-Tools CP-SAT','https://developers.google.com/optimization/cp/cp_solver'),('. Documentação técnica consultada em 27/10/2026.',None)],
      [('[3] EloGroup, atuação institucional','https://elogroup.com/sobre-a-elogroup/'),('; ',None),('CEFET-MG, Engenharia de Computação','https://www.decom.cefetmg.br/'),('; ',None),('orientações do HackatOhm','https://www.eng-eletrica.bh.cefetmg.br/hackatohm/acompanhe/'),('.',None)],
      [('[4] LGPD, Lei 13.709/2018','https://www2.camara.leg.br/legin/fed/lei/2018/lei-13709-14-agosto-2018-787077-normaatualizada-pl.html'),('; ',None),('ANPD, Resolução 19/2024','https://www.gov.br/anpd/pt-br/acesso-a-informacao/institucional/atos-normativos/regulamentacoes_anpd/resolucao-cd-anpd-no-19-de-23-de-agosto-de-2024'),('. Requisitos aplicáveis ao tratamento e à transferência internacional de dados.',None)],
    ]
    for row in rows:
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
        for label,url in row:
            if url: hyperlink(p,label,url)
            else:
                r=p.add_run(label);r.font.size=Pt(8.5)
    p=doc.add_paragraph('Ferramentas de análise e diagramação: Python. Conteúdo revisado sob responsabilidade da equipe.')
    for r in p.runs:r.font.size=Pt(8.5)


def summary():
    doc=new_document()
    title=doc.add_paragraph(TEXT['title'],'Title')
    title.alignment=WD_ALIGN_PARAGRAPH.LEFT
    header=[
      ('Desafio', 'D03 | EloGroup | Gestão inteligente de currículos | Case Nexora'),
      ('Equipe', 'Quatro e meio'),
      ('Curso de todos os integrantes', 'Engenharia de Computação'),
      ('Integrantes', 'Iago Soares Santana — 5º período (líder)\nGabriel Galvão dos Reis Bandeira — 5º período\nBernardo Vieira Rocha — 5º período\nLucas Francischetti Alvarenga — 6º período\nArthur de Oliveira Torres — 5º período'),
      ('Mentores', 'Daniel Hasan Dalip e Poliana Aparecida Corrêa de Oliveira'),
      ('Vídeo', 'https://github.com/santana-iago/hackatohm/releases/tag/pitch-v4'),
    ]
    for label,value in header:
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
        p.add_run(label+': ').bold=True;p.add_run(value)
        for r in p.runs:r.font.size=Pt(9)
    doc.add_paragraph('Síntese','Heading 1')
    body(doc,TEXT['synthesis'],'A contribuição da equipe está na qualidade da implementação proposta: integrar interpretação semântica, regras verificáveis e otimização, preservando a origem de cada decisão.')
    doc.add_paragraph('1  Problema e objetivo','Heading 1');body(doc,TEXT['problem'])
    doc.add_paragraph('2  Solução proposta','Heading 1');body(doc,TEXT['solution'])
    doc.add_paragraph('3  Fundamentação e viabilidade','Heading 1');body(doc,TEXT['foundation'])
    doc.add_paragraph('4  Resultados obtidos e esperados','Heading 1')
    for key,label in [('obtained','Obtidos'),('expected','Esperados')]:
        body(doc,label+': '+TEXT[key],label+':')
    doc.add_page_break()
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(CANVAS),width=Cm(17))
    p=doc.add_paragraph('Figura 1. Arquitetura proposta e implantação gradual. Elaboração da equipe.')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(7.5)
    doc.add_paragraph('5  Afinidade empresa–CEFET-MG e perspectivas de cooperação','Heading 1')
    body(doc,TEXT['partnership'])
    refs(doc)
    doc.core_properties.title=TEXT['title']
    doc.core_properties.subject='Resumo de proposta com arquitetura, viabilidade, impacto, governança e cooperação'
    doc.core_properties.author='Equipe Quatro e meio'
    doc.save(WORD)
    labels = {'synthesis':'Síntese', 'problem':'Problema e objetivo', 'solution':'Solução proposta', 'foundation':'Fundamentação e viabilidade', 'obtained':'Resultados obtidos', 'expected':'Resultados esperados', 'partnership':'Afinidade empresa–CEFET-MG e perspectivas de cooperação'}
    (ROOT/'APTUS_Texto_Resumo_Revisado.md').write_text('\n\n'.join(['# '+TEXT['title']]+['## '+labels[k]+'\n\n'+v for k,v in TEXT.items() if k!='title'])+'\n',encoding='utf-8')
    counts={k:len(v.split()) for k,v in TEXT.items()}
    (ROOT/'word_counts.json').write_text(json.dumps(counts,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(counts,ensure_ascii=False))
    print(WORD)


def pitch():
    VIDEO.mkdir(parents=True, exist_ok=True)
    doc=new_document(10.5)
    doc.add_paragraph('APTUS: roteiro do pitch orientado à avaliação','Title')
    intro='Uso interno da equipe. Meta de montagem: 4min50s, com pausas e transições a conferir no ensaio. Ler somente as falas; as cenas são orientações de gravação. A demonstração pressupõe o mock navegável. Manter as identificações de dados sintéticos e arquitetura proposta.'
    doc.add_paragraph(intro)
    md=['# APTUS: roteiro do pitch orientado à avaliação','',intro,'']
    for i,(time,title,cue,speech) in enumerate(PITCH):
        if i in (3,6):doc.add_page_break()
        doc.add_paragraph(time+'  '+title,'Heading 1')
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
        p.add_run('Cena ').bold=True;p.add_run(cue)
        for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(71,80,96)
        for s in speech.split('\n\n'):body(doc,s)
        md.extend(['## '+time+' '+title,'','**Cena — não ler:** '+cue,'',speech,''])
    doc.add_paragraph('Orientações de ensaio','Heading 1')
    notes=[
      'O pitch cobre as duas rubricas disponíveis: qualidade/viabilidade, problema, impacto, comunicação e diferenciação do briefing; inovação, aderência, viabilidade, impacto e fundamentação da página geral.',
      'O bloco técnico tem 65 segundos previstos. O canvas deve permanecer legível, com destaque progressivo dos grupos. A montagem seleciona tomadas do mock sem acrescentar todo o banco de três minutos ao roteiro.',
      'Os valores de 600–1.000 horas são uma conta sobre as premissas do briefing; 320 horas liberadas são uma hipótese. O exemplo 30/31/33/35 é sintético. Não afirmar economia real, melhora de taxa de vitória ou integração já validada.',
      'O cenário econômico completo do resumo assume R$160/h: 320h × R$160 = R$51.200 de capacidade, antes das despesas adicionais. O benefício precisa superar os custos totais adicionais; capacidade liberada não é caixa garantido.',
      'A evolução de 31 para 35 depende de duas provas distintas, verificadas, aceitas e recalculadas. O teto do cenário é 37, não 35. Falha, ausência de documento e incerteza não autorizam a aprovação.',
      'A redação descreve controles projetados. Benchmark, revisão contratual, testes de segurança e validação de integração ainda são condições de avanço. Não declarar conformidade LGPD concluída.',
    ]
    for note in notes:doc.add_paragraph(note,'List Bullet')
    md.extend(['## Orientações de ensaio','']+['- '+n for n in notes]+[''])
    doc.add_paragraph('Fontes','Heading 1')
    p=doc.add_paragraph('Briefing EloGroup, páginas 14–19; documentação Aptus 3.1; página oficial do HackatOhm 2026; documentação de RAG, Jev e OR-Tools; LGPD e Resolução ANPD 19/2024. As referências completas estão no resumo e na matriz de avaliação.')
    for r in p.runs:r.font.size=Pt(9)
    doc.core_properties.title='APTUS Roteiro do pitch orientado à avaliação'
    doc.save(VIDEO/'APTUS_Roteiro_Pitch_Pipeline_Tecnico.docx')
    (VIDEO/'APTUS_Roteiro_Pitch_Pipeline_Tecnico.md').write_text('\n'.join(md),encoding='utf-8')
    print('Pitch words:',sum(len(x[3].split()) for x in PITCH))

if __name__=='__main__':
    summary()
    pitch()
