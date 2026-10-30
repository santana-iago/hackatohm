import sys, os, json, re
import docx

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
stats = []
for f in sorted(os.listdir(src)):
    p = os.path.join(src, f)
    try:
        d = docx.Document(p)
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                cells = []
                for c in r.cells:
                    if c.text not in cells:
                        cells.append(c.text)
                parts.append(" | ".join(cells))
        txt = "\n".join(x for x in parts if x.strip())
        err = None
    except Exception as e:
        txt, err = "", str(e)
    with open(os.path.join(out, f + ".txt"), "w") as fh:
        fh.write(txt)
    stats.append({"file": f, "chars": len(txt), "words": len(txt.split()), "tables": len(d.tables) if not err else 0, "err": err})
json.dump(stats, open(os.path.join(out, "_stats.json"), "w"), indent=1)
print(len(stats), "files;", sum(1 for s in stats if s["err"]), "errors")
